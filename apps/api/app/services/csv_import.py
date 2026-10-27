"""Spreadsheet (CSV / XLSX) parsing, column mapping suggestions and transactional import."""

import csv
import datetime as dt
import hashlib
import io
import re
import uuid

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import AuditLog, Contact, User
from app.services.contacts import (
    IMPORT_FIELDS,
    apply_fields,
    get_or_create_tags,
    log_activity,
    suppressed_set,
    template_variables,
)
from app.services.normalize import InvalidEmail, normalize_email

# Exact header aliases (after normalisation: lower case, punctuation → spaces).
_ALIASES = {
    "email": ["email", "email address", "e mail", "e mail address", "work email", "emailaddress", "mail",
              "contact email", "owner email", "primary email", "business email"],
    "first_name": ["first name", "firstname", "first", "given name", "fname", "owner first name", "contact first name"],
    "last_name": ["last name", "lastname", "last", "surname", "family name", "lname", "owner last name",
                  "contact last name"],
    # pseudo-field: split into first/last at import time; the original is kept in custom field source_name
    "full_name": ["name", "full name", "contact name", "contact", "client name", "client", "owner name", "owner",
                  "contact person", "agent name", "agent", "bondsman", "principal", "person", "primary contact"],
    "company_name": ["company", "company name", "organization", "organisation", "organization name", "business",
                     "business name", "agency", "agency name", "bail bonds company", "bail bond company",
                     "bail company", "firm", "firm name", "dba", "account name"],
    "website": ["website", "url", "company website", "site", "domain", "web", "web site", "website url", "homepage"],
    "phone": ["phone", "phone number", "telephone", "mobile", "tel", "office phone", "business phone", "cell",
              "cell phone", "main phone", "phone no"],
    "city": ["city", "town"],
    "state": ["state", "region", "province", "st", "state code", "state province"],
    "timezone": ["timezone", "time zone", "tz"],
    "industry": ["industry", "sector", "vertical"],
    "job_title": ["title", "job title", "position", "role"],
    "linkedin_url": ["linkedin", "linkedin url", "linkedin profile"],
    "notes": ["notes", "note", "comments"],
    "source": ["source", "lead source"],
    "tags": ["tags", "tag", "labels"],
}
# Fallback when no exact alias matched: (field, words that must all appear in the header). Order matters.
_KEYWORDS = [
    ("email", ("email",)), ("email", ("e mail",)),
    ("first_name", ("first", "name")), ("last_name", ("last", "name")),
    ("website", ("website",)), ("website", ("url",)),
    ("phone", ("phone",)),
    ("company_name", ("company",)), ("company_name", ("agency",)), ("company_name", ("business",)),
    ("full_name", ("name",)),
]

HONORIFICS = {"mr", "mrs", "ms", "miss", "mx", "dr", "prof", "rev", "hon", "sir", "madam"}
SUFFIXES = {"jr", "sr", "ii", "iii", "iv", "v", "esq", "phd", "md", "cpa", "dds"}
# A "name" containing any of these is a business, not a person: never split it into first/last.
BUSINESS = re.compile(r"\b(llc|l\.l\.c|inc|incorporated|corp|corporation|co|company|bail|bails|bond|bonds|bonding|"
                      r"agency|associates|group|services|enterprises|ltd|lp|llp|pllc|pc|the|office|insurance)\b|[&@\d/]",
                      re.I)
NAME_TOKEN = re.compile(r"[^\W\d_][\w'’.\-]*", re.UNICODE)


class CsvError(ValueError):
    pass


def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", h.strip().lower()).strip()


def _name_headers(headers: list[str]) -> list[str]:
    """Blank headers become "Column N" so every column can still be mapped."""
    out = [h.strip() or f"Column {i}" for i, h in enumerate(headers, start=1)]
    if len(set(out)) != len(out):
        raise CsvError("The file has duplicate column headers")
    return out


def parse_csv(raw: bytes, max_rows: int) -> tuple[list[str], list[list[str]]]:
    if raw.startswith(b"\xef\xbb\xbf"):
        raw = raw[3:]
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    if "\x00" in text:
        raise CsvError("File does not look like a text CSV")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [r for r in reader if any(c.strip() for c in r)]
    if not rows:
        raise CsvError("CSV is empty")
    headers = _name_headers(rows[0])
    body = rows[1:]
    if len(body) > max_rows:
        raise CsvError(f"CSV has {len(body)} rows; the maximum per import is {max_rows}")
    return headers, body


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))  # phone numbers / zips stored as numbers
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return str(v).strip()


MAX_XLSX_UNCOMPRESSED = 200 * 1024 * 1024  # an .xlsx is a zip: cap what it expands to, not just the upload


def parse_xlsx(raw: bytes, max_rows: int) -> tuple[list[str], list[list[str]], str]:
    """First non-empty worksheet. Cached values only (formulas are never evaluated)."""
    import zipfile

    from openpyxl import load_workbook  # uses defusedxml when installed

    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            expanded = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile as exc:
        raise CsvError("Could not read the Excel file. Save it as .xlsx (Excel 2007+) or .csv and try again") from exc
    if expanded > MAX_XLSX_UNCOMPRESSED:
        raise CsvError("This Excel file expands to more than 200 MB. Save it as .csv or split it and try again")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise CsvError("Could not read the Excel file. Save it as .xlsx (Excel 2007+) or .csv and try again") from exc
    try:
        for ws in wb.worksheets:
            rows: list[list[str]] = []
            for r in ws.iter_rows(values_only=True):
                cells = [_cell(v) for v in r]
                if any(cells):
                    rows.append(cells)
                    if len(rows) > max_rows + 1:
                        raise CsvError(f"Sheet “{ws.title}” has more than {max_rows} rows; the maximum per import is {max_rows}")
            if not rows:
                continue
            width = max(len(r) for r in rows)
            while width and not any(len(r) >= width and r[width - 1] for r in rows):
                width -= 1  # drop trailing empty columns
            rows = [(r + [""] * width)[:width] for r in rows]
            return _name_headers(rows[0]), rows[1:], ws.title
    finally:
        wb.close()
    raise CsvError("The workbook has no data")


def parse_upload(filename: str, raw: bytes, max_rows: int) -> tuple[list[str], list[list[str]], dict]:
    name = filename.lower()
    if name.endswith(".xlsx") or raw[:4] == b"PK\x03\x04":
        headers, rows, sheet = parse_xlsx(raw, max_rows)
        return headers, rows, {"format": "xlsx", "sheet": sheet}
    headers, rows = parse_csv(raw, max_rows)
    return headers, rows, {"format": "csv", "sheet": None}


def file_fingerprint(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def previous_import(db: Session, sha256: str) -> dict | None:
    """The last committed import of this exact file, if any (to prevent accidental re-imports)."""
    row = db.scalar(select(AuditLog).where(AuditLog.action == "contacts.imported", AuditLog.data["sha256"].astext == sha256)
                    .order_by(AuditLog.created_at.desc()).limit(1))
    if row is None:
        return None
    return {"at": row.created_at.isoformat(), "filename": row.data.get("filename"), "created": row.data.get("created")}


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    mapping = {h: "" for h in headers}
    used: set[str] = set()
    for h in headers:  # exact aliases first, so "First Name" wins over the generic "name" keyword
        nh = _norm_header(h)
        for field, aliases in _ALIASES.items():
            if field not in used and nh in aliases:
                mapping[h] = field
                used.add(field)
                break
    for h in headers:
        if mapping[h]:
            continue
        nh = _norm_header(h)
        for field, needed in _KEYWORDS:
            if field not in used and all(re.search(rf"\b{re.escape(n)}", nh) for n in needed):
                mapping[h] = field
                used.add(field)
                break
    return mapping


def split_name(raw: str) -> tuple[str | None, str | None, str | None]:
    """(first, last, note). Conservative: anything that isn't clearly a person's name is left unsplit."""
    value = " ".join(raw.split())
    if not value:
        return None, None, None
    if BUSINESS.search(value):
        return None, None, "name_looks_like_business"
    if "," in value:  # "Smith, John A."
        last_part, _, first_part = value.partition(",")
        value = f"{first_part.strip()} {last_part.strip()}".strip()
    tokens = value.split()
    while tokens and tokens[0].rstrip(".").lower() in HONORIFICS:
        tokens.pop(0)
    while tokens and tokens[-1].rstrip(".,").lower() in SUFFIXES:
        tokens.pop()
    if not tokens or not all(NAME_TOKEN.fullmatch(t) for t in tokens):
        return None, None, "name_not_parsed"
    first = tokens[0]
    if len(tokens) == 1:
        return first, None, "single_name"  # never guess a last name
    last = " ".join(t for t in tokens[1:] if not re.fullmatch(r"[^\W\d_]\.?", t))  # drop middle initials
    return first, last or None, "name_split"


def _row_to_data(headers: list[str], row: list[str], mapping: dict[str, str]) -> tuple[dict, list[str]]:
    data: dict = {"custom_fields": {}}
    notes: list[str] = []
    full_name = ""
    for idx, header in enumerate(headers):
        target = mapping.get(header) or ""
        if not target:
            continue
        value = row[idx].strip() if idx < len(row) else ""
        if target.startswith("custom:"):
            key = target[7:]
            if value:
                data["custom_fields"][key] = value[:2000]
        elif target == "tags":
            data["tags"] = [t.strip() for t in re.split(r"[;,|]", value) if t.strip()]
        elif target == "full_name":
            full_name = value
        else:
            data[target] = value
    if full_name:
        data["custom_fields"]["source_name"] = full_name[:300]  # the original is always kept
        first, last, note = split_name(full_name)
        if note:
            notes.append(note)
        if first and not (data.get("first_name") or "").strip():
            data["first_name"] = first
        if last and not (data.get("last_name") or "").strip():
            data["last_name"] = last
    if not data["custom_fields"]:
        data.pop("custom_fields")
    return data, notes


def validate_mapping(headers: list[str], mapping: dict[str, str]) -> None:
    targets = [t for h, t in mapping.items() if t]
    unknown_headers = set(mapping) - set(headers)
    if unknown_headers:
        raise CsvError(f"Mapping references unknown columns: {', '.join(sorted(unknown_headers))}")
    if "email" not in targets:
        raise CsvError("A column must be mapped to email")
    for t in targets:
        if t.startswith("custom:"):
            if not re.fullmatch(r"[a-z][a-z0-9_]{0,49}", t[7:]):
                raise CsvError(f"Invalid custom field name '{t[7:]}'")
        elif t not in IMPORT_FIELDS:
            raise CsvError(f"Unknown target field '{t}'")
    dupes = {t for t in targets if targets.count(t) > 1}
    if dupes:
        raise CsvError(f"Fields mapped more than once: {', '.join(sorted(dupes))}")


REPORT_ROWS = 2000


def import_rows(
    db: Session,
    headers: list[str],
    rows: list[list[str]],
    mapping: dict[str, str],
    *,
    duplicate_strategy: str,
    tags: list[str],
    source: str | None,
    user: User | None,
    dry_run: bool,
) -> tuple[dict, list[uuid.UUID]]:
    """Import inside the caller's transaction. Caller commits (or rolls back for dry runs).

    Every row gets a report entry: new | update | existing_skipped | duplicate_in_file | suppressed |
    invalid_email | missing_email | error, plus warnings (missing_first_name, missing_company_name, name notes).
    """
    validate_mapping(headers, mapping)
    result = {"created": 0, "updated": 0, "skipped_duplicates": 0, "skipped_suppressed": 0, "invalid": 0,
              "missing_fields": 0, "errors": []}
    report: list[dict] = []

    def add(row_no: int, email: str | None, data: dict, status: str, warnings: list[str]) -> None:
        report.append({"row": row_no, "email": email, "first_name": data.get("first_name") or None,
                       "last_name": data.get("last_name") or None, "company_name": data.get("company_name") or None,
                       "source_name": (data.get("custom_fields") or {}).get("source_name"),
                       "status": status, "warnings": warnings})

    prepared: list[tuple[int, str, dict, list[str]]] = []
    seen: set[str] = set()
    for i, row in enumerate(rows, start=2):  # row 1 is the header
        data, notes = _row_to_data(headers, row, mapping)
        raw_email = (data.get("email") or "").strip()
        try:
            email = normalize_email(raw_email)
        except InvalidEmail as exc:
            result["invalid"] += 1
            result["errors"].append({"row": i, "email": raw_email or None, "error": str(exc)})
            add(i, raw_email or None, data, "invalid_email" if raw_email else "missing_email", notes)
            continue
        if email in seen:
            result["skipped_duplicates"] += 1
            result["errors"].append({"row": i, "email": email, "error": "Duplicate email within file"})
            add(i, email, data, "duplicate_in_file", notes)
            continue
        seen.add(email)
        data["email"] = email
        prepared.append((i, email, data, notes))

    emails = [e for _, e, _, _ in prepared]
    suppressed = suppressed_set(db, emails)
    existing = {
        c.email: c
        for c in db.scalars(select(Contact).where(Contact.email.in_(emails), Contact.deleted_at.is_(None)))
    } if emails else {}
    extra_tags = get_or_create_tags(db, tags) if tags else []
    touched: list[uuid.UUID] = []

    for row_no, email, data, notes in prepared:
        if email in suppressed:
            result["skipped_suppressed"] += 1
            add(row_no, email, data, "suppressed", notes)
            continue
        if source and not data.get("source"):
            data["source"] = source
        row_tags = data.pop("tags", None) or []
        contact = existing.get(email)
        if contact is not None and duplicate_strategy == "skip":
            result["skipped_duplicates"] += 1
            add(row_no, email, data, "existing_skipped", notes)
            continue
        try:
            with db.begin_nested():
                if contact is not None:
                    # update only non-empty values so blanks never wipe existing data
                    apply_fields(db, contact, {k: v for k, v in data.items() if v not in ("", None, {})})
                    result["updated"] += 1
                    status = "update"
                else:
                    contact = Contact(email=email, custom_fields={})
                    db.add(contact)
                    apply_fields(db, contact, data)
                    db.flush()
                    log_activity(db, contact.id, "import", "Imported from spreadsheet", user, source=data.get("source"))
                    result["created"] += 1
                    status = "new"
                all_tags = {t.id: t for t in [*contact.tags, *extra_tags, *get_or_create_tags(db, row_tags)]}
                contact.tags = list(all_tags.values())
                db.flush()
                touched.append(contact.id)
        except Exception as exc:  # row-level DB error (e.g. constraint) – report, keep going
            result["errors"].append({"row": row_no, "email": email, "error": f"Database error: {type(exc).__name__}"})
            result["invalid"] += 1
            add(row_no, email, data, "error", notes)
            continue
        # warnings reflect the stored lead (an update may fill gaps from existing data)
        warnings = list(notes)
        merge = template_variables(contact)  # same values {{first_name}} / {{company_name}} will render
        if not merge["first_name"]:
            warnings.append("missing_first_name")
        if not merge["company_name"]:
            warnings.append("missing_company_name")
        if "missing_first_name" in warnings or "missing_company_name" in warnings:
            result["missing_fields"] += 1
        add(row_no, email, {"first_name": contact.first_name, "last_name": contact.last_name,
                            "company_name": merge["company_name"], "custom_fields": data.get("custom_fields")},
            status, warnings)

    report.sort(key=lambda r: r["row"])
    result["errors"] = result["errors"][:500]
    result["rows"] = report[:REPORT_ROWS]
    result["rows_truncated"] = len(report) > REPORT_ROWS
    result["dry_run"] = dry_run
    return result, touched


def export_contacts_csv(contacts: list[Contact]) -> str:
    out = io.StringIO()
    custom_keys = sorted({k for c in contacts for k in (c.custom_fields or {})})
    cols = ["email", "first_name", "last_name", "company_name", "website", "phone", "city", "state", "timezone",
            "industry", "job_title", "linkedin_url", "source", "verification_status", "lead_status", "tags",
            "last_contacted_at", "last_replied_at", "created_at"]
    writer = csv.writer(out)
    writer.writerow(cols + [f"custom:{k}" for k in custom_keys])
    for c in contacts:
        row = []
        for col in cols:
            if col == "tags":
                row.append(";".join(t.name for t in c.tags))
            else:
                v = getattr(c, col)
                row.append("" if v is None else (v.isoformat() if hasattr(v, "isoformat") else str(v)))
        row += [str((c.custom_fields or {}).get(k, "")) for k in custom_keys]
        writer.writerow([_csv_safe(v) for v in row])
    return out.getvalue()


def _csv_safe(value: str) -> str:
    """Neutralise spreadsheet formula injection."""
    if value and value[0] in "=+-@\t\r":
        return "'" + value
    return value
